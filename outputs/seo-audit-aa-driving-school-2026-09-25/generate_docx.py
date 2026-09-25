#!/usr/bin/env python3
"""Build AA-Driving-School-SEO-Audit.docx from AUDIT-REPORT.md + appendices + charts.

Small Markdown converter: headings, paragraphs (**bold**, *italic*, `code`), pipe tables,
bullets, numbered lists, fenced code, images. Appendices A/B come from audit_appendix.py
(data-driven), Appendix C from APPENDIX-C-ORIGINAL-PROMPT.md.

Usage: python3 generate_docx.py
"""
import re, json
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from audit_appendix import add_appendices

OUT = "AA-Driving-School-SEO-Audit.docx"
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def runs(p, text, size=None):
    for part in INLINE.split(text):
        if not part: continue
        if part.startswith("**"): r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`"): r = p.add_run(part[1:-1]); r.font.name = "Consolas"; r.font.color.rgb = RGBColor(0x25, 0x6a, 0xbf)
        elif part.startswith("["): r = p.add_run(re.match(r"\[([^\]]+)\]", part).group(1))
        elif part.startswith("*") and len(part) > 2: r = p.add_run(part[1:-1]); r.italic = True
        else: r = p.add_run(part)
        if size: r.font.size = Pt(size)


def table(doc, lines):
    rows = [[c.strip() for c in l.strip().strip("|").split("|")] for l in lines if not re.match(r"^\|\s*:?-", l.strip())]
    n = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=n); t.style = "Light Grid Accent 1"
    size = 8 if n > 6 else 9
    for i, r in enumerate(rows):
        cells = t.add_row().cells
        for j in range(n):
            cells[j].text = ""
            runs(cells[j].paragraphs[0], r[j] if j < len(r) else "", size)
            if i == 0:
                for rr in cells[j].paragraphs[0].runs: rr.bold = True
    doc.add_paragraph()


def code(doc, lines):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr(); shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), "F3F2EF"); pPr.append(shd)
    r = p.add_run("\n".join(lines)); r.font.name = "Consolas"; r.font.size = Pt(7.5)


def render(doc, md, base_level=0):
    lines = md.split("\n"); i = 0
    while i < len(lines):
        l = lines[i]
        if l.startswith("```"):
            j = i + 1
            while j < len(lines) and not lines[j].startswith("```"): j += 1
            code(doc, lines[i + 1:j]); i = j + 1; continue
        if l.strip().startswith("|"):
            j = i
            while j < len(lines) and lines[j].strip().startswith("|"): j += 1
            table(doc, lines[i:j]); i = j; continue
        m = re.match(r"^(#{1,6}) (.*)", l)
        if m:
            doc.add_heading(re.sub(r"[*`]", "", m.group(2)), min(4, len(m.group(1)) + base_level)); i += 1; continue
        m = re.match(r"^!\[.*?\]\((.*?)\)", l)
        if m:
            doc.add_picture(m.group(1), width=Cm(15.5)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER; i += 1; continue
        m = re.match(r"^(\s*)[-*] (.*)", l)
        if m:
            p = doc.add_paragraph(style="List Bullet 2" if len(m.group(1)) >= 2 else "List Bullet"); runs(p, m.group(2)); i += 1; continue
        m = re.match(r"^\s*\d+\. (.*)", l)
        if m:
            p = doc.add_paragraph(style="List Number"); runs(p, m.group(1)); i += 1; continue
        if l.strip() in ("", "---"): i += 1; continue
        p = doc.add_paragraph(); runs(p, l.strip()); i += 1


def toc(doc):
    p = doc.add_paragraph(); r = p.add_run()
    for tag, txt in (("begin", None), (None, 'TOC \\o "1-2" \\h \\z \\u'), ("separate", None), ("end", None)):
        if tag:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), tag); r._r.append(e)
            if tag == "separate":
                t = OxmlElement("w:t"); t.text = "Right-click and choose Update Field to build the table of contents."; r._r.append(t)
        else:
            e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt; r._r.append(e)


def main():
    scores = json.load(open("scores.json"))
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
    for s in doc.sections: s.left_margin = s.right_margin = Cm(2.2)
    doc.add_paragraph().add_run("\n\n")
    t = doc.add_paragraph(); r = t.add_run("AA Driving School"); r.bold = True; r.font.size = Pt(30)
    t = doc.add_paragraph(); r = t.add_run("Mobile-First SEO Audit: 7 Selected Pages"); r.font.size = Pt(18)
    doc.add_paragraph("aa.co.nz/drivers/driving-school/  ·  25 September 2026")
    t = doc.add_paragraph(); r = t.add_run(f"Overall SEO health: {scores['health']}/100"); r.bold = True; r.font.size = Pt(14)
    doc.add_picture("charts/01_health_gauge.png", width=Cm(9))
    doc.add_paragraph("Data: Playwright rendered DOM (390px), Lighthouse 12.8.2 mobile, Firecrawl NZ SERPs, "
                      "historical SEMrush (June 2026). GBP skipped by instruction. See Appendix A.")
    doc.add_page_break(); doc.add_heading("Contents", 1); toc(doc); doc.add_page_break()

    md = open("AUDIT-REPORT.md").read()
    md = re.sub(r"^# .*\n", "", md, count=1)              # title handled on cover
    md = md.split("\n## Appendices")[0]                    # appendices rendered below
    render(doc, md, base_level=-1)
    appc = add_appendices(doc)
    render(doc, appc, base_level=0)
    doc.save(OUT); print("wrote", OUT)


if __name__ == "__main__":
    main()
